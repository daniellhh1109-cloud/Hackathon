"""Member E: result plots, submission checks and evidence-based deck source."""
import json
from pathlib import Path
import pandas as pd
from src.agent.portfolio_agent import write_json


def build_report(run_dir):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    root=Path(run_dir);status=json.loads((root/'run_result.json').read_text())
    if status['status']!='completed':raise ValueError('report requires a completed backtest')
    output=root/'report'
    t=pd.read_csv(root/'timeline.csv');metrics=json.loads((root/'metrics.json').read_text())
    holdings=pd.read_csv(root/'monthly_holdings.csv')
    if status.get('holdings_weight_unit') != 'percent_of_NAV':
        raise ValueError('Holdings unit is not certified as percent_of_NAV; rerun with the corrected exporter')
    output.mkdir(exist_ok=False)
    x=pd.to_datetime(t.Date)
    for filename,columns,title in [
        ('cumulative.png',['nav','benchmark_nav','market_nav'],'Cumulative NAV (net of declared costs)'),
        ('risk.png',['drawdown'],'Monthly drawdown'),
        ('exposure.png',['gross','net','input_beta'],'Rebalance exposures'),
        ('rolling.png',['rolling_12m_beta','rolling_12m_ir'],'Rolling 12-month beta and active IR')]:
        fig,ax=plt.subplots(figsize=(10,4.7))
        for col in columns:ax.plot(x,t[col],label=col)
        ax.set_title(title);ax.legend();ax.grid(alpha=.2);fig.autofmt_xdate();fig.tight_layout()
        fig.savefig(output/filename,dpi=160);plt.close(fig)
    checked={'monthly_holdings_columns':list(holdings.columns)==['Date','PERMNO','TICKER','COMPANY NAME','WEIGHT'],
             'unique_security_month':not holdings.duplicated(['Date','PERMNO']).any(),
             'real_data':not status.get('synthetic',False),'complete_68_months':status['full_oos'],'historical_names_complete':status['names_complete'],
             'all_months_have_holdings':set(holdings.Date)==set(t.Date),
             'CVs_provided':False,'deck_reviewed':False,'borrow_availability_verified':False}
    write_json(output/'submission_checklist.json',checked)
    t[['Date','total_return']].rename(columns={'total_return':'RETURN'}).to_csv(output/'Portfolio_Returns.csv',index=False)
    holdings.to_csv(output/'Monthly_Holdings.csv',index=False)
    scope=('SYNTHETIC ACCEPTANCE ONLY. ' if status.get('synthetic') else '')+f"{t.Date.iloc[0]} to {t.Date.iloc[-1]}; {len(t)} months; controller={status['controller_kind']}"
    slides=[
        ('Strategy and evaluation scope',[scope,'Shared numerical/text forecast model; constrained long/short allocation.',
          'Partial periods are development evidence, not a complete competition submission.']),
        ('Data and timing',['12 months of 147 factors; optional six-month filing memory.',
          'Annual expanding training and rolling two-year validation.',
          'Current-month realized returns are loaded only after weights are committed.']),
        ('Prediction model',['ModernTCN short-window adaptation; frozen MiniLM in multimodal mode.',
          'Monthly filing attention, temporal attention and gated residual fusion.',f"Zero-benchmark OOS R2: {metrics['oos_r2_zero']}"]),
        ('Agent and optimizer',[f"Controller: {status['controller_kind']}; seven bounded tools.",
          'CVXPY/CLARABEL: score reward, turnover and concentration penalties.',
          'Fixed candidate signs and minimum magnitudes enforce actual position counts.']),
        ('Portfolio returns',[f"CAGR: {metrics['cagr']:.2%}",f"Information ratio: {metrics['information_ratio']}",
          f"Sharpe: {metrics['sharpe']}",'Benchmark: TB3MS / 1200 + 0.04 / 12.']),
        ('Risk and implementation',[f"Maximum monthly drawdown: {metrics['max_drawdown']:.2%}",
          f"Realized beta: {metrics['realized_beta']}; HAC SE: {metrics['beta_se']}",
          f"Mean traded notional: {metrics['mean_traded_notional']:.3f}"]),
        ('Trading assumptions',[status['cash_assumption'],f"Costs: {status['costs']}",
          'Missing held returns cause failure. Borrow availability has not been certified.']),
        ('Reproducibility and limitations',['Versioned configuration, checkpoint provenance, tool logs and holdings hashes.',
          f"Full 68-month period: {status['full_oos']}; historical labels complete: {status['names_complete']}",
          'Complete CVs, historical-name issues and final reviewer checklist before submitting.'])]
    write_json(output/'deck_content.json',[{'title':a,'bullets':b} for a,b in slides])
    lines=['# Eight-slide presentation source',scope,'']
    for i,(title,body) in enumerate(slides,1):lines += [f'## {i}. {title}',*['- '+v for v in body],'']
    (output/'deck_content.md').write_text('\n'.join(lines),encoding='utf-8')
    return checked


def export_pptx(run_dir):
    """Editable eight-slide draft. Export to PDF with PowerPoint/LibreOffice after review."""
    from pptx import Presentation
    from pptx.util import Inches,Pt
    from pptx.dml.color import RGBColor
    root=Path(run_dir)/'report';slides=json.loads((root/'deck_content.json').read_text())
    prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
    for i,item in enumerate(slides):
        slide=prs.slides.add_slide(prs.slide_layouts[6])
        title=slide.shapes.add_textbox(Inches(.65),Inches(.45),Inches(12),Inches(.9)).text_frame
        title.text=item['title'];title.paragraphs[0].font.size=Pt(32);title.paragraphs[0].font.bold=True
        title.paragraphs[0].font.color.rgb=RGBColor.from_string('15334D')
        body=slide.shapes.add_textbox(Inches(.65),Inches(1.65),Inches(12),Inches(4.5)).text_frame
        body.word_wrap=True
        for j,line in enumerate(item['bullets']):
            p=body.paragraphs[0] if j==0 else body.add_paragraph();p.text=line;p.font.size=Pt(22);p.space_after=Pt(20)
        footer=slide.shapes.add_textbox(Inches(.65),Inches(6.8),Inches(12),Inches(.4)).text_frame
        footer.text=f'McGill-FIAM 2026 | Development report | {i+1}/8';footer.paragraphs[0].font.size=Pt(12)
    target=root/'results_draft.pptx'
    if target.exists():raise FileExistsError(target)
    prs.save(target)
    return target
